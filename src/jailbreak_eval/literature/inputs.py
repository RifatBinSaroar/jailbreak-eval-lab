"""Read tabular cells without running formulas or interpreting research prose."""

import csv
import hashlib
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile


class InputError(ValueError):
    """A source needs correction before it can be imported."""


def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).casefold())


def strict_json(text):
    def object_pairs(pairs):
        result = {}
        for name, value in pairs:
            if name in result:
                raise InputError(f"Repeated JSON key: {name}")
            result[name] = value
        return result

    def reject(value):
        raise InputError(f"Non-finite JSON number: {value}")

    return json.loads(text, object_pairs_hook=object_pairs, parse_constant=reject)


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise InputError("Non-finite numeric cell")
    return value


def source(path, sheet, row, columns):
    return {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "sheet": sheet, "row": row, "columns": columns}


def headers(values):
    result = []
    for index, value in enumerate(values, 1):
        # Unnamed columns remain addressable; never silently discard their cells.
        name = str(value) if value is not None and str(value).strip() else f"__column_{index}"
        if key(name) in {key(h) for h in result}:
            raise InputError(f"Repeated/ambiguous column heading: {name!r}")
        result.append(name)
    return result


def read_rows(filename, sheet=None, header_row=1):
    """Return source rows and explicit skipped-sheet notices."""
    path = Path(filename)
    notices, rows = [], []
    suffix = path.suffix.casefold()
    if header_row < 1:
        raise InputError("--header-row must be at least 1")
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise InputError("Install dependencies: python -m pip install -r requirements-literature.txt") from exc
        try:
            workbook = load_workbook(path, data_only=False, keep_links=False)
        except (BadZipFile, KeyError) as exc:
            raise InputError(f"Invalid XLSX workbook: {path.name}") from exc
        try:
            if sheet and sheet not in workbook.sheetnames:
                raise InputError(f"Sheet {sheet!r} not found. Available: {', '.join(workbook.sheetnames)}")
            for tab in workbook:
                if sheet and tab.title != sheet:
                    continue
                names = headers([c.value for c in tab[header_row]])
                if not sheet and not ({"paper", "title", "papertitle"} & {key(h) for h in names}):
                    notices.append(f"Skipped sheet {tab.title!r}: no Paper/Title heading")
                    continue
                for cells in tab.iter_rows(min_row=header_row + 1):
                    if not any(c.value is not None or c.hyperlink for c in cells):
                        continue
                    columns = []
                    for name, cell in zip(names, cells):
                        columns.append({"header": name, "value": scalar(cell.value),
                                        "cell": cell.coordinate,
                                        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                                        "is_formula": cell.data_type == "f"})
                    rows.append(source(path, tab.title, cells[0].row, columns))
        finally:
            workbook.close()
    elif suffix in {".csv", ".tsv"}:
        if sheet:
            raise InputError("--sheet is only supported for XLSX")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            table = list(csv.reader(handle, delimiter="\t" if suffix == ".tsv" else ","))
        if len(table) < header_row:
            raise InputError("Header row is outside the file")
        names = headers(table[header_row - 1])
        for number, values in enumerate(table[header_row:], header_row + 1):
            if not any(v.strip() for v in values):
                continue
            if len(values) > len(names):
                raise InputError(f"Row {number} has more cells than headings")
            values += [None] * (len(names) - len(values))
            columns = [{"header": h, "value": v, "cell": str(i), "hyperlink": None,
                        "is_formula": isinstance(v, str) and v.startswith("=")}
                       for i, (h, v) in enumerate(zip(names, values), 1)]
            rows.append(source(path, None, number, columns))
    elif suffix in {".json", ".jsonl"}:
        if sheet or header_row != 1:
            raise InputError("--sheet/--header-row do not apply to structured JSON notes")
        content = path.read_text(encoding="utf-8-sig")
        if suffix == ".jsonl":
            items = [(i, strict_json(line)) for i, line in enumerate(content.splitlines(), 1) if line.strip()]
        else:
            data = strict_json(content)
            if not isinstance(data, list):
                raise InputError("JSON input must be an array of paper objects (use export for a registry)")
            items = list(enumerate(data, 1))
        for number, item in items:
            if not isinstance(item, dict) or not item:
                raise InputError(f"JSON record {number} must be a nonempty object")
            headers(item.keys())
            columns = [{"header": h, "value": v, "cell": h, "hyperlink": None, "is_formula": False}
                       for h, v in item.items()]
            rows.append(source(path, None, number, columns))
    else:
        raise InputError("Use .xlsx, .csv, .tsv, .json or .jsonl. Save older .xls files as .xlsx in Excel.")
    if not rows:
        raise InputError("No paper rows found. Check --sheet and --header-row.")
    return rows, notices
